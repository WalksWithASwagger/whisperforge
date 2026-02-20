"""
Content Export Utilities
========================

Export generated content to various file formats (Text, Markdown, Word, PDF).
"""

from datetime import datetime
from io import BytesIO


def create_text_export(results: dict) -> str:
    """Create a formatted plain-text export of all content."""
    lines = [
        "=" * 60,
        "WHISPERFORGE CONTENT EXPORT",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "=" * 60,
        "",
    ]

    sections = [
        ("AUDIO TRANSCRIPT", results.get("transcript", "")),
        ("EXTRACTED WISDOM", results.get("wisdom", "")),
        ("CONTENT OUTLINE", results.get("outline", "")),
        ("FULL ARTICLE", results.get("article", "")),
        ("SOCIAL MEDIA CONTENT", results.get("social_content", "")),
    ]

    for title, content in sections:
        if content:
            lines.extend([f"## {title}", "-" * 40, content, "", ""])

    if results.get("notion_url"):
        lines.extend(["## NOTION LINK", "-" * 40, results["notion_url"], ""])

    return "\n".join(lines)


def export_to_markdown(results: dict) -> str:
    """Export results to Markdown format."""
    lines = ["# WhisperForge Content Export"]
    sections = [
        ("## Transcript", results.get("transcript", "")),
        ("## Wisdom", results.get("wisdom", "")),
        ("## Outline", results.get("outline", "")),
        ("## Article", results.get("article", "")),
        ("## Social Content", results.get("social_content", "")),
    ]
    for title, content in sections:
        if content:
            lines.extend([title, "", content, ""])
    if results.get("notion_url"):
        lines.extend(["## Notion Link", results["notion_url"]])
    return "\n".join(lines)


def export_to_word(results: dict) -> bytes:
    """Export results to a Word document (.docx)."""
    from docx import Document

    doc = Document()
    doc.add_heading("WhisperForge Content Export", level=1)
    for title, content in [
        ("Transcript", results.get("transcript", "")),
        ("Wisdom", results.get("wisdom", "")),
        ("Outline", results.get("outline", "")),
        ("Article", results.get("article", "")),
        ("Social Content", results.get("social_content", "")),
    ]:
        if content:
            doc.add_heading(title, level=2)
            doc.add_paragraph(content)
    if results.get("notion_url"):
        doc.add_heading("Notion Link", level=2)
        doc.add_paragraph(results["notion_url"])

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()


def export_to_pdf(results: dict) -> bytes:
    """Export results to a PDF file."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)
    text = create_text_export(results)
    for line in text.split("\n"):
        pdf.cell(0, 10, txt=line, ln=1)
    return pdf.output(dest="S").encode("latin-1")
